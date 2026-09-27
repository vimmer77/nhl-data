#!/usr/bin/env python3
"""
Daily NHL data build for the Bet365 moneyline model (and future player props).

Runs on GitHub Actions (open internet). Pulls:
  - NHL stats API (api.nhle.com/stats/rest): per-game skater and goalie logs, team PP/PK, prior-season aggregates
  - NHL web API (api-web.nhle.com): schedules, rosters, standings
  - MoneyPuck season summaries: team 5v5 xG/high-danger/PDO, goalie GSAx

Writes (in output/):
  nhl_daily_data.xlsx   README / Schedule / Goalies / Teams / Skaters tabs
  *.csv                 same tables as CSV
  meta.json             data-through date, build time, season info

Per-game logs are stored incrementally in data/<seasonId>/ so each run only fetches recent days.
"""
import datetime as dt
import io
import os
import json
import math
import pathlib
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from zoneinfo import ZoneInfo

import numpy as np
import pandas as pd

ROOT = pathlib.Path(__file__).resolve().parent.parent
DATA = ROOT / os.environ.get("DATA_DIR", "data")
OUT = ROOT / os.environ.get("OUT_DIR", "output")
PT = ZoneInfo("America/Los_Angeles")
UA = {"User-Agent": "Mozilla/5.0 (personal NHL research; github.com/vimmer77/nhl-data)"}
STATS = "https://api.nhle.com/stats/rest/en"
WEB = "https://api-web.nhle.com/v1"
MP = "https://moneypuck.com/moneypuck/playerData/seasonSummary"

# ---- rules (kept in one place so they match the project instructions) ----
GOALIE_PRIOR_FULL_UNTIL = 10      # starts; last season's GSAx weight fades to 0 by this many starts
GOALIE_SHRINK_MIN = 600           # phantom league-average minutes added to every goalie's rate
GOALIE_QUALIFY_MIN = 540          # pooled minutes needed to be ranked among "regular workload" goalies
ROLE_MIN_GP_CURRENT = 5           # skater games before current-season usage replaces last season's
TREND_MIN_GP = 8                  # skater games this season before a hot/cold label is given
TREND_WINDOW = 10
TREND_POINT_GAP = 3               # points above/below expectation over the window
TREND_RATIO = 1.4                 # and at least this ratio (hot) / its inverse (cold)


# ---------------------------------------------------------------- fetching
def fetch(url, kind="json", retries=4, allow_404=False):
    last = None
    for i in range(retries):
        try:
            with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=90) as r:
                b = r.read()
            return json.loads(b) if kind == "json" else b.decode("utf-8")
        except urllib.error.HTTPError as e:
            if e.code == 404 and allow_404:
                return None
            last = e
        except Exception as e:  # network hiccup
            last = e
        time.sleep(2 * (i + 1))
    raise RuntimeError(f"fetch failed: {url} -> {last!r}")


def stats_query(path, cayenne, is_game, is_agg=False, sort=None):
    q = {"isAggregate": str(is_agg).lower(), "isGame": str(is_game).lower(),
         "start": 0, "limit": -1, "cayenneExp": cayenne}
    if sort:
        q["sort"] = json.dumps(sort)
    d = fetch(f"{STATS}/{path}?{urllib.parse.urlencode(q)}")
    rows, total = d.get("data", []), d.get("total", 0)
    if total and len(rows) < total:  # limit=-1 capped: page through
        rows, start = [], 0
        while start < total:
            q.update(start=start, limit=100)
            rows += fetch(f"{STATS}/{path}?{urllib.parse.urlencode(q)}").get("data", [])
            start += 100
    return rows


def moneypuck(year, kind, cache_ok):
    """Season summary CSV. Prior seasons are cached; the current one is refreshed daily."""
    f = DATA / "moneypuck" / f"{year}_{kind}.csv"
    f.parent.mkdir(parents=True, exist_ok=True)
    if cache_ok and f.exists():
        return pd.read_csv(f)
    txt = fetch(f"{MP}/{year}/regular/{kind}.csv", kind="text", allow_404=True)
    if txt is None:
        return None
    f.write_text(txt)
    return pd.read_csv(io.StringIO(txt))


# ---------------------------------------------------------------- season setup
def season_ids(today):
    start = today.year if today.month >= 7 else today.year - 1
    return start, f"{start}{start + 1}", f"{start - 1}{start}"


def team_list():
    s = fetch(f"{WEB}/standings/now")
    return sorted({t["teamAbbrev"]["default"] for t in s["standings"]})


def club_schedules(teams, sid):
    games = {}
    per_team = {}
    for t in teams:
        d = fetch(f"{WEB}/club-schedule-season/{t}/{sid}", allow_404=True) or {}
        dates = []
        for g in d.get("games", []):
            if g.get("gameType") != 2:
                continue
            gd = g["gameDate"]
            dates.append(gd)
            games[g["id"]] = {"gameId": g["id"], "date": gd, "startTimeUTC": g.get("startTimeUTC"),
                              "away": g["awayTeam"]["abbrev"], "home": g["homeTeam"]["abbrev"],
                              "state": g.get("gameState")}
        per_team[t] = sorted(set(dates))
    return pd.DataFrame(games.values()), per_team


def rosters(teams):
    rows = []
    for t in teams:
        d = fetch(f"{WEB}/roster/{t}/current", allow_404=True) or {}
        for grp in ("forwards", "defensemen", "goalies"):
            for p in d.get(grp, []):
                rows.append({"playerId": p["id"], "team": t, "pos": p.get("positionCode"),
                             "name": f'{p["firstName"]["default"]} {p["lastName"]["default"]}'})
    return pd.DataFrame(rows)


# ---------------------------------------------------------------- game logs (incremental)
def update_logs(sid, start_year, today):
    folder = DATA / sid
    folder.mkdir(parents=True, exist_ok=True)
    fs, fg = folder / "skater_games.csv", folder / "goalie_games.csv"
    sk = pd.read_csv(fs) if fs.exists() else pd.DataFrame()
    gl = pd.read_csv(fg) if fg.exists() else pd.DataFrame()
    since = f"{start_year}-09-01"
    if len(sk):
        since = (pd.to_datetime(sk["gameDate"]).max() - pd.Timedelta(days=3)).strftime("%Y-%m-%d")
    until = (dt.date.fromisoformat(today) - dt.timedelta(days=1)).isoformat()  # completed games only
    srt = [{"property": "gameDate", "direction": "ASC"}, {"property": "playerId", "direction": "ASC"}]

    def windowed(path):  # the API silently caps a response at 10,000 rows, so ask 10 days at a time
        rows, d0, end = [], dt.date.fromisoformat(since), dt.date.fromisoformat(until)
        while d0 <= end:
            d1 = min(d0 + dt.timedelta(days=9), end)
            cay = f'seasonId={sid} and gameTypeId=2 and gameDate>="{d0}" and gameDate<="{d1} 23:59:59"'
            got = stats_query(path, cay, True, sort=srt)
            if len(got) >= 10000:
                raise RuntimeError(f"{path} window {d0}..{d1} hit the 10,000-row cap")
            rows += got
            d0 = d1 + dt.timedelta(days=1)
        return pd.DataFrame(rows)

    summ = windowed("skater/summary")
    toi = windowed("skater/timeonice")
    if len(summ):
        keep = ["playerId", "gameId", "gameDate", "skaterFullName", "teamAbbrev", "opponentTeamAbbrev",
                "homeRoad", "positionCode", "goals", "assists", "points", "shots", "ppPoints"]
        new = summ[keep].merge(toi[["playerId", "gameId", "timeOnIce", "ppTimeOnIce"]],
                               on=["playerId", "gameId"], how="left")
        sk = pd.concat([sk, new]).drop_duplicates(["playerId", "gameId"], keep="last")
        sk.to_csv(fs, index=False)

    gsum = windowed("goalie/summary")
    if len(gsum):
        keep = ["playerId", "gameId", "gameDate", "goalieFullName", "teamAbbrev", "opponentTeamAbbrev",
                "gamesStarted", "shotsAgainst", "saves", "goalsAgainst", "timeOnIce"]
        gl = pd.concat([gl, gsum[keep]]).drop_duplicates(["playerId", "gameId"], keep="last")
        gl.to_csv(fg, index=False)
    return sk, gl


def prior_skater_agg(prior_sid):
    f = DATA / prior_sid / "skater_agg.csv"
    if f.exists():
        return pd.read_csv(f)
    f.parent.mkdir(parents=True, exist_ok=True)
    cay = f"seasonId={prior_sid} and gameTypeId=2"
    s = pd.DataFrame(stats_query("skater/summary", cay, False, is_agg=True))
    t = pd.DataFrame(stats_query("skater/timeonice", cay, False, is_agg=True))
    df = s[["playerId", "skaterFullName", "positionCode", "gamesPlayed", "goals", "assists", "points",
            "shots", "timeOnIcePerGame"]].merge(t[["playerId", "ppTimeOnIcePerGame"]], on="playerId", how="left")
    df.to_csv(f, index=False)
    return df


def team_summary(sid, id2abbr):
    rows = stats_query("team/summary", f"seasonId={sid} and gameTypeId=2", False)
    df = pd.DataFrame(rows)
    if not len(df):
        return pd.DataFrame(columns=["team", "gp", "w", "l", "otl", "pp", "pk"])
    df["team"] = df["teamId"].map(id2abbr)
    return pd.DataFrame({"team": df["team"], "gp": df["gamesPlayed"], "w": df["wins"], "l": df["losses"],
                         "otl": df["otLosses"], "pp": df["powerPlayPct"] * 100, "pk": df["penaltyKillPct"] * 100})


# ---------------------------------------------------------------- goalies
def goalie_table(mp_cur, mp_prior, gl, ros):
    def mp_all(df):
        if df is None or not len(df):
            return pd.DataFrame(columns=["playerId", "mp_name", "mp_team", "gp", "toi_min", "gsax"])
        a = df[df["situation"] == "all"].assign(gsax=lambda x: x["xGoals"] - x["goals"])
        a = a.groupby("playerId", as_index=False).agg(mp_name=("name", "last"), mp_team=("team", "last"),
                                                      gp=("games_played", "sum"), toi_sec=("icetime", "sum"),
                                                      gsax=("gsax", "sum"))
        return a.assign(toi_min=a["toi_sec"] / 60).drop(columns="toi_sec")

    cur, pri = mp_all(mp_cur), mp_all(mp_prior)
    ids = set(cur["playerId"]) | set(pri["playerId"]) | set(ros.loc[ros["pos"] == "G", "playerId"])
    if len(gl):
        ids |= set(gl["playerId"])
    df = pd.DataFrame({"playerId": sorted(ids)})
    df = df.merge(cur.add_suffix("_cur").rename(columns={"playerId_cur": "playerId"}), on="playerId", how="left")
    df = df.merge(pri.add_suffix("_pri").rename(columns={"playerId_pri": "playerId"}), on="playerId", how="left")

    if len(gl):
        starts = gl.groupby("playerId")["gamesStarted"].sum()
        last = gl[gl["gamesStarted"] == 1].groupby("playerId")["gameDate"].max()
        lteam = gl.sort_values("gameDate").groupby("playerId")["teamAbbrev"].last()
        gname = gl.groupby("playerId")["goalieFullName"].last()
        team_starts = gl[gl["gamesStarted"] == 1].groupby("teamAbbrev").size()
    else:
        starts = last = lteam = gname = pd.Series(dtype=object)
        team_starts = pd.Series(dtype=float)
    df["starts"] = df["playerId"].map(starts).fillna(0).astype(int)
    df["last_start"] = df["playerId"].map(last)
    r = ros[ros["pos"] == "G"].set_index("playerId")
    df["team"] = df["playerId"].map(r["team"]).fillna(df["playerId"].map(lteam)).fillna(df["mp_team_cur"]).fillna(df["mp_team_pri"])
    df["name"] = df["playerId"].map(r["name"]).fillna(df["playerId"].map(gname)).fillna(df["mp_name_cur"]).fillna(df["mp_name_pri"])
    # keep goalies who are on a roster now, or started this season
    df = df[df["playerId"].isin(r.index) | (df["starts"] > 0)].copy()

    for c in ["gsax_cur", "toi_min_cur", "gsax_pri", "toi_min_pri", "gp_pri", "gp_cur"]:
        df[c] = df[c].fillna(0.0)
    df["w_prior"] = (1 - df["starts"] / GOALIE_PRIOR_FULL_UNTIL).clip(lower=0)
    num = df["gsax_cur"] + df["w_prior"] * df["gsax_pri"]
    den = df["toi_min_cur"] + df["w_prior"] * df["toi_min_pri"]
    df["pooled_min"] = den
    df["gsax60"] = num / ((den + GOALIE_SHRINK_MIN) / 60)
    df["team_start_share"] = np.where(df["team"].map(team_starts).fillna(0) > 0,
                                      df["starts"] / df["team"].map(team_starts).replace(0, np.nan), np.nan)

    q = df[df["pooled_min"] >= GOALIE_QUALIFY_MIN].sort_values("gsax60", ascending=False).copy()
    n = len(q)
    q["rank"] = np.arange(1, n + 1)
    top3, mid3 = math.ceil(n / 3), math.ceil(2 * n / 3)
    q["tier"] = np.select([q["rank"] <= 5, q["rank"] <= top3, q["rank"] <= mid3],
                          ["Elite", "Above-average", "Average"], "Replacement-level")
    avg_cut = q.loc[q["rank"] == mid3, "gsax60"].min() if n else 0
    df = df.merge(q[["playerId", "rank", "tier"]], on="playerId", how="left")
    nodata = (df["toi_min_cur"] + df["toi_min_pri"]) == 0
    small = df["tier"].isna()
    df.loc[small & (df["gsax60"] >= avg_cut) & ~nodata, "tier"] = "Average (small sample)"
    df.loc[small & (df["gsax60"] < avg_cut) & ~nodata, "tier"] = "Replacement-level (small sample)"
    df.loc[small & nodata, "tier"] = "Replacement-level (no NHL data)"

    out = pd.DataFrame({
        "Goalie": df["name"], "Team": df["team"], "Tier": df["tier"], "Rank": df["rank"],
        "GSAx/60 (blended)": df["gsax60"].round(3),
        "Starts this season": df["starts"], "GSAx this season": df["gsax_cur"].round(1),
        "GP last season": df["gp_pri"].astype(int), "GSAx last season": df["gsax_pri"].round(1),
        "Last-season weight": df["w_prior"].round(2),
        "Team start share": (df["team_start_share"] * 100).round(0),
        "Last start": df["last_start"], "playerId": df["playerId"],
    })
    order = {"Elite": 0, "Above-average": 1, "Average": 2, "Average (small sample)": 3,
             "Replacement-level": 4, "Replacement-level (small sample)": 5, "Replacement-level (no NHL data)": 6}
    out["_o"] = out["Tier"].map(order)
    return out.sort_values(["Team", "_o", "Starts this season"], ascending=[True, True, False]).drop(columns="_o")


# ---------------------------------------------------------------- teams
def team_game_rows(sk):
    if not len(sk):
        return pd.DataFrame(columns=["team", "gameId", "gameDate", "gf", "sf", "ga", "sa"])
    t = sk.groupby(["teamAbbrev", "gameId", "gameDate", "opponentTeamAbbrev"], as_index=False)[["goals", "shots"]].sum()
    t = t.rename(columns={"teamAbbrev": "team", "goals": "gf", "shots": "sf"})
    opp = t[["team", "gameId", "gf", "sf"]].rename(columns={"team": "opponentTeamAbbrev", "gf": "ga", "sf": "sa"})
    return t.merge(opp, on=["opponentTeamAbbrev", "gameId"], how="left")


def mp_team5(df):
    if df is None or not len(df):
        return pd.DataFrame(columns=["team", "xgf", "hdcf", "pdo", "gp5"])
    a = df[df["situation"] == "5on5"]
    xf, xa = a["flurryScoreVenueAdjustedxGoalsFor"], a["flurryScoreVenueAdjustedxGoalsAgainst"]
    hf, ha = a["highDangerShotsFor"], a["highDangerShotsAgainst"]
    sh = a["goalsFor"] / a["shotsOnGoalFor"]
    sv = 1 - a["goalsAgainst"] / a["shotsOnGoalAgainst"]
    return pd.DataFrame({"team": a["team"], "xgf": 100 * xf / (xf + xa), "hdcf": 100 * hf / (hf + ha),
                         "pdo": sh + sv, "gp5": a["games_played"]})


def team_table(teams, ts_cur, ts_pri, mp_cur, mp_pri, sk, gl):
    df = pd.DataFrame({"team": teams})
    df = df.merge(ts_cur, on="team", how="left").merge(
        ts_pri[["team", "pp", "pk"]].rename(columns={"pp": "pp_pri", "pk": "pk_pri"}), on="team", how="left")
    c5, p5 = mp_team5(mp_cur), mp_team5(mp_pri)
    df = df.merge(c5.drop(columns="gp5"), on="team", how="left")
    df = df.merge(p5.drop(columns="gp5").add_suffix("_pri").rename(columns={"team_pri": "team"}), on="team", how="left")

    tg = team_game_rows(sk)
    last10 = tg.sort_values("gameDate").groupby("team").tail(10).groupby("team")[["gf", "ga", "sf", "sa"]].sum()
    df["l10_gf"] = df["team"].map(last10["gf"]) if len(last10) else np.nan
    df["l10_ga"] = df["team"].map(last10["ga"]) if len(last10) else np.nan
    df["l10_sog"] = df["team"].map(100 * last10["sf"] / (last10["sf"] + last10["sa"])) if len(last10) else np.nan
    if len(gl):
        st = gl[gl["gamesStarted"] == 1].sort_values("gameDate", ascending=False)
        recent = st.groupby("teamAbbrev")["goalieFullName"].apply(lambda s: " / ".join(s.head(5)))
        df["recent_starters"] = df["team"].map(recent)
    else:
        df["recent_starters"] = ""
    df["gp"] = df["gp"].fillna(0).astype(int)
    rec = df.apply(lambda r: f'{int(r.w)}-{int(r.l)}-{int(r.otl)}' if r.gp > 0 else "", axis=1)
    return pd.DataFrame({
        "Team": df["team"], "GP": df["gp"], "Record (W-L-OTL)": rec,
        "PP% this season": df["pp"].round(1), "PK% this season": df["pk"].round(1),
        "PP% last season": df["pp_pri"].round(1), "PK% last season": df["pk_pri"].round(1),
        "5v5 xGF% this season": df["xgf"].round(1), "5v5 HDCF% this season": df["hdcf"].round(1),
        "5v5 PDO this season": df["pdo"].round(3),
        "5v5 xGF% last season": df["xgf_pri"].round(1), "5v5 HDCF% last season": df["hdcf_pri"].round(1),
        "5v5 PDO last season": df["pdo_pri"].round(3),
        "Last 10: GF": df["l10_gf"], "Last 10: GA": df["l10_ga"], "Last 10: shot share %": df["l10_sog"].round(1),
        "Recent starting goalies (newest first)": df["recent_starters"],
    })


# ---------------------------------------------------------------- schedule / rest
def schedule_table(games, per_team, gl, gp_by_team, today, days=3):
    if not len(games):
        return pd.DataFrame()
    want = [(today + dt.timedelta(days=i)).isoformat() for i in range(days)]
    g = games[games["date"].isin(want)].sort_values(["date", "startTimeUTC"])
    last_starter = {}
    if len(gl):
        st = gl[gl["gamesStarted"] == 1].sort_values("gameDate")
        last_starter = st.groupby("teamAbbrev")["goalieFullName"].last().to_dict()

    def rest(team, date):
        ds = [d for d in per_team.get(team, []) if d < date]
        if not ds:
            return "Season opener", "No", "No"
        d0 = dt.date.fromisoformat(date)
        prev = dt.date.fromisoformat(ds[-1])
        r = (d0 - prev).days - 1
        in3 = sum(1 for d in ds if (d0 - dt.date.fromisoformat(d)).days <= 3)
        return r, "Yes" if r == 0 else "No", "Yes" if in3 >= 2 else "No"

    rows = []
    for x in g.itertuples():
        start = pd.Timestamp(x.startTimeUTC)
        ar, ab, a3 = rest(x.away, x.date)
        hr, hb, h3 = rest(x.home, x.date)
        rows.append({"Date": x.date, "Start (ET)": start.tz_convert("America/New_York").strftime("%-I:%M %p"),
                     "Start (PT)": start.tz_convert(PT).strftime("%-I:%M %p"), "Away": x.away, "Home": x.home,
                     "Away rest days": ar, "Away back-to-back": ab, "Away 3-in-4": a3,
                     "Home rest days": hr, "Home back-to-back": hb, "Home 3-in-4": h3,
                     "Away GP": gp_by_team.get(x.away, 0), "Home GP": gp_by_team.get(x.home, 0),
                     "Away last starting goalie": last_starter.get(x.away, ""),
                     "Home last starting goalie": last_starter.get(x.home, "")})
    return pd.DataFrame(rows)


# ---------------------------------------------------------------- skaters
def skater_table(sk, prior, ros):
    ros_s = ros[ros["pos"] != "G"].copy()
    if len(sk):
        sk = sk.sort_values("gameDate")
        tg = sk.groupby("teamAbbrev")["gameId"].unique()
        team_last10 = {t: set(pd.Series(ids).tail(10)) for t, ids in tg.items()}
        cur = sk.groupby("playerId").agg(gp=("gameId", "nunique"), g=("goals", "sum"), a=("assists", "sum"),
                                         p=("points", "sum"), s=("shots", "sum"), toi=("timeOnIce", "sum"),
                                         pptoi=("ppTimeOnIce", "sum"), name=("skaterFullName", "last"),
                                         pos=("positionCode", "last"), team=("teamAbbrev", "last"))
        recent = sk[[gid in team_last10.get(t, ()) for t, gid in zip(sk["teamAbbrev"], sk["gameId"])]]
        recent_players = recent[["playerId", "teamAbbrev"]].drop_duplicates()
        l10gp = recent.groupby("playerId")["gameId"].nunique()
    else:
        cur = pd.DataFrame(columns=["gp", "g", "a", "p", "s", "toi", "pptoi", "name", "pos", "team"])
        recent_players = pd.DataFrame(columns=["playerId", "teamAbbrev"])
        l10gp = pd.Series(dtype=float)

    base = ros_s[["playerId", "team", "pos", "name"]]
    extra = recent_players.rename(columns={"teamAbbrev": "team"})
    extra = extra[~extra["playerId"].isin(base["playerId"])]
    if len(extra):
        extra = extra.assign(pos=extra["playerId"].map(cur["pos"]), name=extra["playerId"].map(cur["name"]))
    df = pd.concat([base, extra]).drop_duplicates("playerId")
    df = df.join(cur.drop(columns=["name", "pos", "team"]), on="playerId")
    pr = prior.set_index("playerId")
    df["p_gp"] = df["playerId"].map(pr["gamesPlayed"]).fillna(0)
    df["p_toi_pg"] = df["playerId"].map(pr["timeOnIcePerGame"]).fillna(0) / 60
    df["p_pp_pg"] = df["playerId"].map(pr["ppTimeOnIcePerGame"]).fillna(0) / 60
    for c in ["gp", "g", "a", "p", "s", "toi", "pptoi"]:
        df[c] = df[c].fillna(0)
    df["toi_pg"] = np.where(df["gp"] > 0, df["toi"] / 60 / df["gp"].replace(0, np.nan), np.nan)
    df["pp_pg"] = np.where(df["gp"] > 0, df["pptoi"] / 60 / df["gp"].replace(0, np.nan), np.nan)

    # last-N window per player
    win = {}
    if len(sk):
        for pid, grp in sk.groupby("playerId"):
            w = grp.tail(TREND_WINDOW)
            win[pid] = dict(n=len(w), g=w["goals"].sum(), a=w["assists"].sum(), p=w["points"].sum(),
                            s=w["shots"].sum(), toi=w["timeOnIce"].sum() / 60, pp=w["ppTimeOnIce"].sum() / 60)
    W = pd.DataFrame.from_dict(win, orient="index")
    for c in ["n", "g", "a", "p", "s", "toi", "pp"]:
        df[f"w_{c}"] = df["playerId"].map(W[c]) if len(W) else np.nan
    df["l10_team_gp"] = df["playerId"].map(l10gp).fillna(0)

    # usage basis for roles: recent games if enough, else this season, else last season
    use_recent = df["w_n"].fillna(0) >= 3
    use_cur = (df["gp"] >= ROLE_MIN_GP_CURRENT) & ~use_recent
    df["basis_toi"] = np.select([use_recent, use_cur], [df["w_toi"] / df["w_n"], df["toi_pg"]], df["p_toi_pg"])
    df["basis_pp"] = np.select([use_recent, use_cur], [df["w_pp"] / df["w_n"], df["pp_pg"]], df["p_pp_pg"])
    df["basis_src"] = np.select([use_recent, use_cur], ["last games", "this season"], "last season")
    df.loc[(df["basis_src"] == "last season") & (df["p_gp"] == 0), "basis_src"] = "no NHL data"

    df["Role"] = "Depth F"
    df["PP1"] = ""
    for t, grp in df.groupby("team"):
        d = grp[grp["pos"] == "D"].sort_values("basis_toi", ascending=False)
        for i, idx in enumerate(d.index):
            df.loc[idx, "Role"] = "Top-pair D" if i < 2 else ("Top-4 D" if i < 4 else "Depth D")
        f = grp[grp["pos"] != "D"].sort_values("basis_toi", ascending=False)
        for i, idx in enumerate(f.index):
            df.loc[idx, "Role"] = "Top-6 F" if i < 6 else "Depth F"
        c = f[f["pos"] == "C"]
        if len(c):
            df.loc[c.index[0], "Role"] = "Top-line C"
        pp = grp[grp["basis_pp"] >= 1.0].sort_values("basis_pp", ascending=False).head(5)
        df.loc[pp.index, "PP1"] = "Yes"
    df.loc[df["basis_src"] == "no NHL data", "Role"] = "Depth (no NHL data)"

    # hot / cold
    labels, detail = [], []
    for r in df.itertuples():
        if r.gp < TREND_MIN_GP or not r.w_n or r.w_n < TREND_WINDOW / 2:
            labels.append("Not enough games"); detail.append(""); continue
        wprior = max(0.0, 1 - (r.gp - r.w_n) / 30)
        p_toi_total = r.p_toi_pg * r.p_gp
        pri = pr.loc[r.playerId] if r.playerId in pr.index else None
        p_pts = pri["points"] if pri is not None else 0
        p_g = pri["goals"] if pri is not None else 0
        p_s = pri["shots"] if pri is not None else 0
        b_toi = (r.toi / 60 - r.w_toi) + wprior * p_toi_total
        b_pts = (r.p - r.w_p) + wprior * p_pts
        b_g = (r.g - r.w_g) + wprior * p_g
        b_s = (r.s - r.w_s) + wprior * p_s
        b_gp = (r.gp - r.w_n) + wprior * r.p_gp
        b_pp_total = (r.pptoi / 60 - r.w_pp) + wprior * r.p_pp_pg * r.p_gp
        if b_toi < 300 or b_gp < 5:
            labels.append("Not enough baseline"); detail.append(""); continue
        exp_p = b_pts / b_toi * r.w_toi
        shots_ratio = (r.w_s / r.w_toi) / (b_s / b_toi) if b_s > 0 else np.nan
        sh_pct = b_g / b_s if b_s > 0 else 0.1
        goal_luck = r.w_g - sh_pct * r.w_s
        toi_d = r.w_toi / r.w_n - b_toi / b_gp
        pp_d = r.w_pp / r.w_n - b_pp_total / b_gp
        info = (f"{int(r.w_p)} pts vs {exp_p:.1f} expected in last {int(r.w_n)}; TOI {toi_d:+.1f} min/gm; "
                f"PP {pp_d:+.1f} min/gm; shots rate x{shots_ratio:.2f}; goals vs shooting% {goal_luck:+.1f}")
        hot = (r.w_p - exp_p >= TREND_POINT_GAP) and (r.w_p >= TREND_RATIO * exp_p)
        cold = (exp_p - r.w_p >= TREND_POINT_GAP) and (r.w_p <= exp_p / TREND_RATIO)
        if hot:
            lab = ("Hot – usage-backed" if (toi_d >= 1.5 or pp_d >= 1.0 or shots_ratio >= 1.25)
                   else "Hot – shooting luck" if goal_luck >= 2 else "Hot – mixed")
        elif cold:
            lab = ("Cold – usage down" if (toi_d <= -1.5 or pp_d <= -1.0)
                   else "Cold – shots steady" if shots_ratio >= 0.8 else "Cold – shots down")
        else:
            lab = "Neutral"
        labels.append(lab); detail.append(info)
    df["Trend"], df["Trend detail"] = labels, detail

    ro = {"Top-line C": 0, "Top-pair D": 1, "Top-6 F": 2, "Top-4 D": 3, "Depth F": 4, "Depth D": 5}
    df["_o"] = df["Role"].map(ro).fillna(9)
    out = pd.DataFrame({
        "Player": df["name"], "Team": df["team"], "Pos": df["pos"], "Role": df["Role"], "PP1": df["PP1"],
        "Role based on": df["basis_src"], "Usage TOI/GP": df["basis_toi"].round(1), "Usage PP TOI/GP": df["basis_pp"].round(1),
        "GP": df["gp"].astype(int), "G": df["g"].astype(int), "A": df["a"].astype(int), "P": df["p"].astype(int),
        "SOG": df["s"].astype(int), "TOI/GP": df["toi_pg"].round(1), "PP TOI/GP": df["pp_pg"].round(1),
        "Last 10 GP": df["w_n"].fillna(0).astype(int), "Last 10 G": df["w_g"].fillna(0).astype(int),
        "Last 10 P": df["w_p"].fillna(0).astype(int), "Last 10 SOG": df["w_s"].fillna(0).astype(int),
        "Trend": df["Trend"], "Trend detail": df["Trend detail"],
        "GP last season": df["p_gp"].astype(int), "TOI/GP last season": df["p_toi_pg"].round(1),
        "playerId": df["playerId"], "_o": df["_o"], "_t": df["basis_toi"],
    })
    return out.sort_values(["Team", "_o", "_t"], ascending=[True, True, False]).drop(columns=["_o", "_t"])


# ---------------------------------------------------------------- output
README_NOTES = [
    ("What this is", "Daily data for the NHL Bet365 moneyline project. Attach this file with your screenshots."),
    ("Goalies tab", "Tier comes from blended GSAx per 60 (MoneyPuck). Last season's GSAx fades out over a goalie's first "
                    f"{GOALIE_PRIOR_FULL_UNTIL} starts, and every rate is pulled toward average by {GOALIE_SHRINK_MIN} "
                    "minutes of league-average play so a few hot games can't make someone Elite. Ranked goalies "
                    f"(at least {GOALIE_QUALIFY_MIN} weighted minutes): top 5 Elite, rest of top third Above-average, "
                    "middle third Average, bottom third Replacement-level. Others are marked small sample."),
    ("Teams tab", "5v5 xGF% is MoneyPuck's score- and venue-adjusted expected-goals share. Context only: these "
                  "numbers never move a probability in the model."),
    ("Schedule tab", "Today and the next two days. Rest days: 0 = second night of a back-to-back. "
                     "'Last starting goalie' helps judge who the market expects in net."),
    ("Skaters tab", "Role comes from ice time (recent games when available): top-line C, top-pair D, top-6 F, "
                    "top-4 D, depth. PP1 = top 5 on the team by power-play time. Trend compares the last 10 games "
                    "with the player's baseline; labels explain whether a streak is backed by more ice time/shots "
                    "or by shooting luck. Trends are context only for the moneyline model."),
    ("Sources", "NHL stats API and NHL web API (api.nhle.com, api-web.nhle.com); MoneyPuck season summaries "
                "(moneypuck.com/data.htm)."),
]


def write_outputs(meta, sheets):
    OUT.mkdir(exist_ok=True)
    for name, df in sheets.items():
        df.to_csv(OUT / f"{name.lower()}.csv", index=False)
    (OUT / "meta.json").write_text(json.dumps(meta, indent=1))
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    path = OUT / "nhl_daily_data.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        info = [("Data through", meta["data_through"]), ("Built (Pacific)", meta["built_pt"]),
                ("Season", meta["season"]), ("Regular-season games played so far", meta["games_played"]),
                ("Next scheduled game date", meta.get("next_game_date") or "")] + [("", "")] + README_NOTES
        pd.DataFrame(info, columns=["Item", "Value"]).to_excel(xw, sheet_name="README", index=False)
        for name, df in sheets.items():
            df.drop(columns=[c for c in df.columns if c == "playerId"]).to_excel(xw, sheet_name=name, index=False)
        head = PatternFill("solid", fgColor="1F3864")
        for ws in xw.book.worksheets:
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = head
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.freeze_panes = "A2" if ws.title == "README" else "C2"
            if ws.title != "README":
                ws.auto_filter.ref = ws.dimensions
            for i, col in enumerate(ws.columns, 1):
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[get_column_letter(i)].width = min(max(8, width + 2), 60 if ws.title != "README" else 110)
            if ws.title == "README":
                for row in ws.iter_rows(min_row=2):
                    for c in row:
                        c.alignment = Alignment(wrap_text=True, vertical="top")


def main():
    now_pt = dt.datetime.now(PT)
    if os.environ.get("BUILD_DATE"):  # test mode: pretend it's this date (logs only up to the day before)
        now_pt = dt.datetime.combine(dt.date.fromisoformat(os.environ["BUILD_DATE"]), dt.time(7, 0), PT)
    today = now_pt.date()
    start_year, sid, prior_sid = season_ids(today)
    print(f"build {now_pt:%Y-%m-%d %H:%M} PT, season {sid}, prior {prior_sid}")

    teams = team_list()
    id2abbr = {t["id"]: t["triCode"] for t in fetch(f"{STATS}/team")["data"]}
    games, per_team = club_schedules(teams, sid)
    ros = rosters(teams)
    sk, gl = update_logs(sid, start_year, today.isoformat())
    prior = prior_skater_agg(prior_sid)
    ts_cur, ts_pri = team_summary(sid, id2abbr), team_summary(prior_sid, id2abbr)
    mp_cur_t = moneypuck(start_year, "teams", cache_ok=False)
    mp_pri_t = moneypuck(start_year - 1, "teams", cache_ok=True)
    mp_cur_g = moneypuck(start_year, "goalies", cache_ok=False)
    mp_pri_g = moneypuck(start_year - 1, "goalies", cache_ok=True)
    print(f"teams {len(teams)}, sched games {len(games)}, roster rows {len(ros)}, skater-game rows {len(sk)}, "
          f"goalie-game rows {len(gl)}, prior skaters {len(prior)}, MoneyPuck current {'yes' if mp_cur_t is not None else 'not yet'}")

    goalies = goalie_table(mp_cur_g, mp_pri_g, gl, ros)
    team_tab = team_table(teams, ts_cur, ts_pri, mp_cur_t, mp_pri_t, sk, gl)
    gp_by_team = dict(zip(team_tab["Team"], team_tab["GP"]))
    sched = schedule_table(games, per_team, gl, gp_by_team, today)
    skaters = skater_table(sk, prior, ros)

    played = int(sk["gameId"].nunique()) if len(sk) else 0
    through = str(pd.to_datetime(sk["gameDate"]).max().date()) if len(sk) else "No regular-season games yet"
    future = games[games["date"] >= today.isoformat()]["date"] if len(games) else pd.Series(dtype=str)
    meta = {"data_through": through, "built_pt": f"{now_pt:%Y-%m-%d %H:%M}", "season": sid,
            "games_played": played, "next_game_date": future.min() if len(future) else None,
            "moneypuck_current_season": mp_cur_t is not None}
    write_outputs(meta, {"Schedule": sched, "Goalies": goalies, "Teams": team_tab, "Skaters": skaters})
    print(json.dumps(meta))
    print(f"rows: schedule {len(sched)}, goalies {len(goalies)}, teams {len(team_tab)}, skaters {len(skaters)}")


if __name__ == "__main__":
    main()
